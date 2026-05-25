import PyPDF2
import docx
import io
import time
import logging

logger = logging.getLogger(__name__)

def parse_document(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from PDF, DOCX, or TXT."""
    ext = filename.lower().split(".")[-1]
    t0 = time.perf_counter()

    if ext == "pdf":
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        pages = len(reader.pages)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        text = text.strip()
        logger.info(f"PDF parsed | pages={pages} | chars={len(text)} | elapsed={time.perf_counter()-t0:.3f}s")
        return text

    elif ext == "docx":
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs]).strip()
        logger.info(f"DOCX parsed | paragraphs={len(doc.paragraphs)} | chars={len(text)} | elapsed={time.perf_counter()-t0:.3f}s")
        return text

    elif ext == "txt":
        text = file_bytes.decode("utf-8").strip()
        logger.info(f"TXT parsed | chars={len(text)} | elapsed={time.perf_counter()-t0:.3f}s")
        return text

    else:
        raise ValueError(f"Unsupported file type: {ext}")
